import os
import fitz  # PyMuPDF
import docx

def create_test_files():
    fixtures_dir = os.path.join(os.path.dirname(__file__), "fixtures")
    os.makedirs(fixtures_dir, exist_ok=True)
    
    # 1. Create English Operating Systems PDF with Deadlock definition
    pdf_path = os.path.join(fixtures_dir, "os_deadlock_sample.pdf")
    doc = fitz.open()
    
    # Page 1
    page1 = doc.new_page()
    page1.insert_text(
        (50, 72),
        "Chapter 7: Deadlocks in Operating Systems\n\n"
        "In a multiprogramming environment, several processes may compete for a finite number of resources.\n"
        "Deadlock occurs when multiple processes wait indefinitely for resources held by each other.\n"
        "A deadlock situation can arise if and only if four conditions hold simultaneously:\n"
        "1. Mutual Exclusion: At least one resource must be held in a non-shareable mode.\n"
        "2. Hold and Wait: A process must be currently holding at least one resource and requesting additional resources.\n"
        "3. No Preemption: Resources cannot be preempted from a process holding them.\n"
        "4. Circular Wait: A closed chain of processes exists such that each process holds at least one resource needed by the next process.",
        fontsize=11
    )
    
    # Page 2
    page2 = doc.new_page()
    page2.insert_text(
        (50, 72),
        "Deadlock Handling Strategies\n\n"
        "Operating systems typically handle deadlocks using three main methods:\n"
        "1. Deadlock Prevention: Designing the system so that one of the four necessary conditions cannot occur.\n"
        "2. Deadlock Avoidance: Banker's Algorithm is used to dynamically ensure the system never enters an unsafe state.\n"
        "3. Deadlock Detection and Recovery: Allowing the system to enter a deadlock state, detecting it, and aborting processes to recover.",
        fontsize=11
    )
    doc.save(pdf_path)
    doc.close()
    print(f"Created PDF fixture: {pdf_path}")
    
    # 2. Create Hindi DBMS sample DOCX
    docx_path = os.path.join(fixtures_dir, "dbms_hindi_sample.docx")
    d = docx.Document()
    d.add_heading("डेटाबेस मैनेजमेंट सिस्टम (DBMS) नोट्स", level=1)
    d.add_paragraph(
        "डेटाबेस मैनेजमेंट सिस्टम एक सॉफ्टवेयर है जिसका उपयोग डेटा को व्यवस्थित, संगृहीत और पुनः प्राप्त करने के लिए किया जाता है।"
    )
    d.add_paragraph(
        "रिलेशनल डेटाबेस में डेटा को टेबल (संबंधों) के रूप में स्टोर किया जाता है जिसमें पंक्तियाँ (rows) और कॉलम (columns) होते हैं।"
    )
    d.add_paragraph(
        "प्राइमरी की (Primary Key) प्रत्येक रिकॉर्ड की विशिष्ट पहचान करने के लिए उपयोग की जाती है। यह कभी भी NULL नहीं हो सकती।"
    )
    d.save(docx_path)
    print(f"Created DOCX fixture: {docx_path}")

if __name__ == "__main__":
    create_test_files()
